from pathlib import Path

from openpyxl import load_workbook

from ai_pipeline.sheets import write_workbook


def test_workbook_has_required_tabs(tmp_path: Path):
    output = tmp_path / "pipeline.xlsx"
    write_workbook({}, output)
    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["Startups", "Products", "Research Papers", "Jobs", "News", "Entity Mapping Log"]